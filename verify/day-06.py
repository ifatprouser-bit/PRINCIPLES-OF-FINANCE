#!/usr/bin/env python3
"""
verify/day-06.py
Recomputes every number that appears on day-06.html: the recall box, all four
teaching panels (worked examples and try-its), the quiz answers and the quiz
distractors, plus the coordinates of the price-against-yield curve drawn in the
panel 2 SVG.

Run:  python3 verify/day-06.py
Each line prints a label, the computed value, the value written on the page,
and PASS or FAIL. Exit code is 1 if anything failed.
"""

import math
import os
import sys

FAILS = []
COUNT = 0


def chk(label, computed, on_page, places=2):
    """Compare a computed number with what the page says, at `places` decimals."""
    global COUNT
    COUNT += 1
    c = round(float(computed) + 0.0, places)
    p = round(float(on_page) + 0.0, places)
    ok = abs(c - p) < 10 ** (-places) / 2
    if not ok:
        FAILS.append(label)
    print("%-62s computed %14.6f   page %14.6f   %s"
          % (label, computed, on_page, "PASS" if ok else "FAIL"))


def pv_annuity(pmt, r, n):
    return pmt * (1 - (1 + r) ** (-n)) / r


def pv_single(cf, r, n):
    return cf / (1 + r) ** n


def bond_price(par, coupon_rate, ytm, years):
    return pv_annuity(par * coupon_rate, ytm, years) + pv_single(par, ytm, years)


print("=" * 118)
print("DAY 6  Bonds, and the risks inside them")
print("=" * 118)

# ------------------------------------------------------------------ #
# 0. Source check: the course Excel model, recomputed from scratch.
#    Par 1,000 / 20 years / coupon 6% / required return 5%.
#    The file itself is NOT trusted; the numbers below are computed here.
# ------------------------------------------------------------------ #
print("\n--- SOURCE CHECK: course model 'Annuity_PV_FV_AND COUPON BOND VALUE.xlsm' ---")
XLSM = ("/root/.claude/projects/-home-claude/9644a518-dd62-530e-89e8-35b5a6441735/"
        "tool-results/project-file-940f17ff-a50c-4b35-b25e-f3d8f845d1f1-"
        "Annuity_PV_FV_AND_COUPON_BOND_VALUE.xlsm")
own_price = bond_price(1000, 0.06, 0.05, 20)
print("independently recomputed price of the course model bond: %.4f" % own_price)
if os.path.exists(XLSM):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(XLSM, data_only=True)
        ws = wb["Coupon bond"]
        print("value stored in the workbook (cell C6):              %.4f" % ws["C6"].value)
        print("agreement with our own computation:                  %s"
              % ("yes" if abs(ws["C6"].value - own_price) < 0.01 else "NO"))
    except Exception as exc:                                    # pragma: no cover
        print("could not open the workbook (%s). Page uses our own number." % exc)
else:
    print("workbook not present on this machine. Page uses our own number.")

# ------------------------------------------------------------------ #
# 1. Recall box (day 2: a single cash flow)
# ------------------------------------------------------------------ #
print("\n--- RECALL BOX (day 2) ---")
chk("recall PV: 5,000 at end of year 4, r = 7%", pv_single(5000, 0.07, 4), 3814.48)
chk("recall FV: 5,000 today for 4 years at 7%", 5000 * 1.07 ** 4, 6553.98)
chk("recall growth factor 1.07^4", 1.07 ** 4, 1.310796, 6)
chk("recall check: PV grown back 4 years", pv_single(5000, 0.07, 4) * 1.07 ** 4, 5000.00)

# ------------------------------------------------------------------ #
# 2. PANEL 1  Value the bond
# ------------------------------------------------------------------ #
print("\n--- PANEL 1: value the bond ---")
chk("panel 1 coupon payment: 6% of 1,000", 0.06 * 1000, 60.00)
chk("panel 1 final year cash flow: coupon + par", 0.06 * 1000 + 1000, 1060.00)

chk("worked ex 1 growth factor 1.05^20", 1.05 ** 20, 2.653298, 6)
chk("worked ex 1 annuity factor (1-1.05^-20)/0.05",
    (1 - 1.05 ** -20) / 0.05, 12.462210, 6)
pv_c = pv_annuity(60, 0.05, 20)
pv_f = pv_single(1000, 0.05, 20)
chk("worked ex 1 PV of the 20 coupons of 60", pv_c, 747.73)
chk("worked ex 1 PV of the 1,000 par value", pv_f, 376.89)
chk("worked ex 1 bond price (premium bond)", pv_c + pv_f, 1124.62)
chk("worked ex 1 check: price above par by", pv_c + pv_f - 1000, 124.62)

chk("panel 1 perpetuity bond: 50 forever at 8%", 50 / 0.08, 625.00)

chk("try-it 1A zero coupon 1,000 in 6 years at 6%", pv_single(1000, 0.06, 6), 704.96)
chk("try-it 1A growth factor 1.06^6", 1.06 ** 6, 1.418519, 6)
chk("try-it 1B coupon payment: 4.5% of 5,000", 0.045 * 5000, 225.00)
chk("try-it 1B final year cash flow", 0.045 * 5000 + 5000, 5225.00)

# ------------------------------------------------------------------ #
# 3. PANEL 2  Price against rates
# ------------------------------------------------------------------ #
print("\n--- PANEL 2: price against rates ---")
chk("worked ex 2 coupon payment: 4% of 1,000", 0.04 * 1000, 40.00)
chk("worked ex 2 growth factor 1.06^10", 1.06 ** 10, 1.790848, 6)
chk("worked ex 2 annuity factor (1-1.06^-10)/0.06",
    (1 - 1.06 ** -10) / 0.06, 7.360087, 6)
pv_c2 = pv_annuity(40, 0.06, 10)
pv_f2 = pv_single(1000, 0.06, 10)
chk("worked ex 2 PV of the 10 coupons of 40", pv_c2, 294.40)
chk("worked ex 2 PV of the 1,000 par value", pv_f2, 558.39)
chk("worked ex 2 bond price (discount bond)", pv_c2 + pv_f2, 852.80)
chk("worked ex 2 check: price below par by", 1000 - (pv_c2 + pv_f2), 147.20)

chk("try-it 2 same bond at 4% (coupon 6%, 5 years)", bond_price(1000, 0.06, 0.04, 5), 1089.04)
chk("try-it 2 same bond at 6% (coupon 6%, 5 years)", bond_price(1000, 0.06, 0.06, 5), 1000.00)
chk("try-it 2 same bond at 8% (coupon 6%, 5 years)", bond_price(1000, 0.06, 0.08, 5), 920.15)

# the curve drawn in the panel 2 SVG: price of that same 5 year 6% bond
print("\n--- PANEL 2 SVG: price-against-yield curve, plotted points ---")
SVG_X0, SVG_X1 = 70.0, 500.0          # yield axis, 2% at x0 to 10% at x1
SVG_YTOP, SVG_YBOT = 40.0, 190.0      # price axis, 1,200 at top to 840 at bottom
P_TOP, P_BOT = 1200.0, 840.0
PAGE_POINTS = [                       # copied out of the polyline in the page
    (70.0, 44.8), (123.8, 66.1), (177.5, 86.2), (231.2, 105.3),
    (285.0, 123.3), (338.8, 140.4), (392.5, 156.6), (446.2, 172.0),
    (500.0, 186.5),
]
for i, ypct in enumerate([2, 3, 4, 5, 6, 7, 8, 9, 10]):
    price = bond_price(1000, 0.06, ypct / 100.0, 5)
    x = SVG_X0 + (SVG_X1 - SVG_X0) * (ypct - 2) / 8.0
    y = SVG_YTOP + (SVG_YBOT - SVG_YTOP) * (P_TOP - price) / (P_TOP - P_BOT)
    chk("svg point yield %2d%%  x" % ypct, x, PAGE_POINTS[i][0], 1)
    chk("svg point yield %2d%%  y (price %8.2f)" % (ypct, price), y, PAGE_POINTS[i][1], 1)
par_y = SVG_YTOP + (SVG_YBOT - SVG_YTOP) * (P_TOP - 1000.0) / (P_TOP - P_BOT)
chk("svg par line y for price 1,000", par_y, 123.3, 1)
# the five price labels printed on that drawing, rounded to whole currency
chk("svg label top axis, price at yield 2%", round(bond_price(1000, 0.06, 0.02, 5)), 1189)
chk("svg label bottom axis, price at yield 10%", round(bond_price(1000, 0.06, 0.10, 5)), 848)
chk("svg label 'premium', price at yield 4%", round(bond_price(1000, 0.06, 0.04, 5)), 1089)
chk("svg label 'at par', price at yield 6%", round(bond_price(1000, 0.06, 0.06, 5)), 1000)
chk("svg label 'discount', price at yield 8%", round(bond_price(1000, 0.06, 0.08, 5)), 920)

# ------------------------------------------------------------------ #
# 4. PANEL 3  The government ladder
# ------------------------------------------------------------------ #
print("\n--- PANEL 3: T-bills, T-notes, T-bonds ---")
chk("worked ex 3 yield on a 1-year bill bought at 95.60",
    (100 / 95.60 - 1) * 100, 4.60)
chk("worked ex 3 ratio 100 / 95.60", 100 / 95.60, 1.046025, 6)
chk("worked ex 3 the gain in money terms", 100 - 95.60, 4.40)
chk("worked ex 3 price if buyers now require 6%", 100 / 1.06, 94.34)
chk("worked ex 3 price fall in money terms", 95.60 - 100 / 1.06, 1.26)
chk("try-it 3 price of a 1-year bill at 3.5%", 100 / 1.035, 96.62)
chk("try-it 3 price of the same bill at 5%", 100 / 1.05, 95.24)

# ------------------------------------------------------------------ #
# 5. PANEL 4  The risks
# ------------------------------------------------------------------ #
print("\n--- PANEL 4: risk-free against risky ---")
chk("worked ex 4 government 1-year zero at 4%", pv_single(1000, 0.04, 1), 961.54)
chk("worked ex 4 company 1-year zero at 7%", pv_single(1000, 0.07, 1), 934.58)
chk("worked ex 4 price gap between the two", pv_single(1000, 0.04, 1) - pv_single(1000, 0.07, 1), 26.96)
chk("worked ex 4 credit spread in percentage points", (0.07 - 0.04) * 100, 3.00)
chk("worked ex 4 real return, course approximation 4% - 3%", 4.0 - 3.0, 1.00)
chk("try-it 4 bank: paid to insured depositors", min(105, 100), 100.00)
chk("try-it 4 bank: paid to subordinated debtholders", min(max(105 - 100, 0), 20), 5.00)
chk("try-it 4 bank: loss carried by the subordinated debtholders",
    20 - min(max(105 - 100, 0), 20), 15.00)
chk("try-it 4 bank: paid to shareholders", max(105 - 100 - 20, 0), 0.00)
chk("try-it 4 bank: the three payments add back to the assets",
    min(105, 100) + min(max(105 - 100, 0), 20) + max(105 - 100 - 20, 0), 105.00)

# ------------------------------------------------------------------ #
# 6. QUIZ
# ------------------------------------------------------------------ #
print("\n--- QUIZ ---")
# Q1 concept: coupon 70 on par 1,000, buyers require 5%
chk("Q1 coupon rate implied by a 70 coupon on 1,000 par", 70 / 1000 * 100, 7.00)
chk("Q1 the yield it is compared against", 5.00, 5.00)
chk("Q1 explanation: same bond over 10 years at 5% (a premium price)",
    bond_price(1000, 0.07, 0.05, 10), 1154.43)

# Q2 par 1,000, coupon 5%, 4 years, yield 9%
q2 = bond_price(1000, 0.05, 0.09, 4)
chk("Q2 ANSWER price (coupon 5%, yield 9%, 4 years)", q2, 870.41)
chk("Q2 growth factor 1.09^4", 1.09 ** 4, 1.411582, 6)
chk("Q2 annuity factor (1-1.09^-4)/0.09", (1 - 1.09 ** -4) / 0.09, 3.239720, 6)
chk("Q2 PV of the four coupons of 50", pv_annuity(50, 0.09, 4), 161.99)
chk("Q2 PV of the par value at 9%", pv_single(1000, 0.09, 4), 708.43)
chk("Q2 distractor: discounted at the coupon rate 5% instead of the yield",
    bond_price(1000, 0.05, 0.05, 4), 1000.00)
chk("Q2 distractor: coupons only, par forgotten", pv_annuity(50, 0.09, 4), 161.99)
chk("Q2 distractor: coupons at 9% but par at the coupon rate 5%",
    pv_annuity(50, 0.09, 4) + pv_single(1000, 0.05, 4), 984.69)

# Q3 zero coupon, 1,000 in 3 years, priced 816.30
chk("Q3 the price quoted in the stem (1,000 at 7% for 3 years)",
    pv_single(1000, 0.07, 3), 816.30)
chk("Q3 explanation: face divided by price", 1000 / 816.30, 1.225040, 6)
chk("Q3 explanation: cube root of that ratio", (1000 / 816.30) ** (1 / 3.0), 1.0700, 4)
chk("Q3 ANSWER yield from price 816.30", ((1000 / 816.30) ** (1 / 3.0) - 1) * 100, 7.00)
chk("Q3 distractor: total gain spread evenly over the price",
    ((1000 - 816.30) / 816.30) / 3 * 100, 7.50)
chk("Q3 distractor: total gain spread evenly over the face value",
    ((1000 - 816.30) / 1000) / 3 * 100, 6.12)
chk("Q3 distractor: whole-period return, never annualised",
    (1000 - 816.30) / 816.30 * 100, 22.50)

# Q4 one-year bill at 96.15
chk("Q4 yield on a 1-year bill priced 96.15", (100 / 96.15 - 1) * 100, 4.00)
chk("Q4 what the price becomes if buyers require 5%", 100 / 1.05, 95.24)

# Q5 perpetuity bond, 45 forever at 6%
chk("Q5 ANSWER perpetuity bond 45 forever at 6%", 45 / 0.06, 750.00)
chk("Q5 distractor: value pushed forward one year", 45 / 0.06 * 1.06, 795.00)
chk("Q5 distractor: treated as one payment at year 1", 45 / 1.06, 42.45)
chk("Q5 distractor: assumed par of 1,000", 1000.00, 1000.00)

# Q6 concept: corporate against government, no arithmetic
chk("Q6 concept item carries no arithmetic", 0, 0)

# ------------------------------------------------------------------ #
print("\n" + "=" * 118)
print("numbers checked: %d" % COUNT)
if FAILS:
    print("FAILED: %d" % len(FAILS))
    for f in FAILS:
        print("   - " + f)
    sys.exit(1)
print("ALL PASSED")
